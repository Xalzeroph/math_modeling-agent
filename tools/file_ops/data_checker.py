#!/usr/bin/env python3
"""数据读取规范器 — 确保确定性数据读取，产出格式化 JSON"""

import argparse
import json
import sys
from pathlib import Path
from typing import Any
import re


def _resolve_root() -> Path:
    d = Path.cwd()
    for _ in range(5):
        if (d / "algorithms").exists():
            return d
        d = d.parent
    return Path.cwd()


def check_excel(filepath: Path) -> dict:
    """检查 Excel 文件"""
    try:
        import openpyxl
    except ImportError:
        return {"status": "error", "message": "需要 openpyxl"}

    wb = openpyxl.load_workbook(filepath, data_only=True)
    result = {"sheets": wb.sheetnames, "files": []}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            result["files"].append({"sheet": sheet_name, "rows": 0, "cols": 0})
            continue

        header = rows[0]
        data_rows = rows[1:]
        n_rows = len(data_rows)
        n_cols = len(header)

        # 统计缺失值
        missing = 0
        for row in data_rows:
            for cell in row:
                if cell is None or (isinstance(cell, str) and cell.strip() == ""):
                    missing += 1

        # 列类型推断
        col_types = {}
        for i, col_name in enumerate(header):
            values = [row[i] for row in data_rows if i < len(row) and row[i] is not None]
            if not values:
                col_types[str(col_name)] = "empty"
            else:
                types = set(type(v).__name__ for v in values)
                col_types[str(col_name)] = ", ".join(types)

        result["files"].append({
            "sheet": sheet_name,
            "rows": n_rows,
            "cols": n_cols,
            "missing_cells": missing,
            "col_types": col_types,
        })

    wb.close()
    return {"status": "ok", "type": "xlsx", **result}


def check_csv(filepath: Path) -> dict:
    """检查 CSV 文件 — 自动检测编码"""
    import csv

    # 尝试多种编码
    for enc in ["utf-8", "utf-8-sig", "gbk", "gb2312", "gb18030", "latin-1"]:
        try:
            with open(filepath, "r", encoding=enc) as f:
                reader = csv.reader(f)
                header = next(reader)
                rows = [row for row in reader]
                break
        except (UnicodeDecodeError, StopIteration):
            continue
    else:
        return {"status": "error", "message": "无法识别文件编码"}

    n_rows = len(rows)
    n_cols = len(header)

    missing = sum(1 for row in rows for cell in row if not cell.strip())
    col_types = {}
    for i, col_name in enumerate(header):
        values = [row[i] for row in rows if i < len(row) and row[i].strip()]
        if values:
            types = set()
            for v in values:
                try:
                    float(v)
                    types.add("numeric")
                except ValueError:
                    types.add("string")
            col_types[str(col_name)] = ", ".join(types)

    return {
        "status": "ok",
        "type": "csv",
        "encoding": enc,
        "rows": n_rows,
        "cols": n_cols,
        "missing_cells": missing,
        "col_types": col_types,
    }


def main():
    parser = argparse.ArgumentParser(description="数据读取规范器 — 确定性编码检测和格式报告")
    sub = parser.add_subparsers(dest="action", required=True)

    sub.add_parser("list", help="列出 data/ 中的所有文件")

    info = sub.add_parser("info", help="打印数据概况")
    info.add_argument("--file", required=True, help="数据文件路径")

    validate = sub.add_parser("validate", help="检测数据质量")
    validate.add_argument("--file", required=True, help="数据文件路径")

    args = parser.parse_args()
    root = _resolve_root()

    if args.action == "list":
        data_dir = root / "data"
        if not data_dir.exists():
            data_dir = root / "sessions"
        if not data_dir.exists():
            print(json.dumps({"status": "error", "message": "data/ 和 sessions/ 都不存在"}))
            sys.exit(1)
        files = list(data_dir.rglob("*"))
        result = {
            "status": "ok",
            "files": [{
                "name": str(f.relative_to(data_dir)),
                "size": f.stat().st_size,
                "type": f.suffix.lower(),
            } for f in files if f.suffix.lower() in (".csv", ".xlsx", ".xls", ".txt", ".json")]
        }
        print(json.dumps(result, indent=2, ensure_ascii=False))
        return

    filepath = Path(args.file)
    if not filepath.is_absolute():
        filepath = root / filepath

    if not filepath.exists():
        print(json.dumps({"status": "error", "message": f"文件不存在: {args.file}"}))
        sys.exit(1)

    suffix = filepath.suffix.lower()
    if suffix in (".xlsx", ".xls"):
        result = check_excel(filepath)
    elif suffix == ".csv":
        result = check_csv(filepath)
    else:
        result = {"status": "ok", "type": suffix, "size": filepath.stat().st_size}

    if args.action == "validate":
        result["action"] = "validate"
        issues = []
        if result.get("missing_cells", 0) > 0:
            issues.append(f"缺失值: {result['missing_cells']} 个")
        col_types = result.get("col_types", {})
        for col, t in col_types.items():
            if "string" in t and "numeric" in t:
                issues.append(f"列 '{col}' 类型不统一: {t}")

        if issues:
            result["issues"] = issues

    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
